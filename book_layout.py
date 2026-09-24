"""Workbook layout helpers shared by the three books.

  key_numbers(wb, fin, sym_map)   a compact "Key numbers" column (P/B · EV/EBITDA
                                  · net debt/mcap · 52w position) inserted right
                                  after each tab's Name column -- the reader sees
                                  the financials next to the name, not off-screen
  strength(wb, populations)       ONE ranking scale: "Strength %ile" next to the
                                  name on every thesis / signal tab = the name's
                                  percentile among EVERYTHING that layer scored
                                  (100 = strongest), not just the rows shown
  regroup(wb, groups, colors)     reorder tabs by how a reader works (Decide ·
                                  Theses · Signals · Evidence · Plumbing), colour
                                  the tabs by group, rewrite Contents with a group
                                  column
  tear_sheets(wb, fin, names, ...) one block per name: identity, financial panel,
                                  and every tab the name appears on with its
                                  strength percentile and that tab's own facts
"""

from __future__ import annotations

import json
from bisect import bisect_left, bisect_right
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path("/home/user/cyclepapa")
NAVY = "1F3864"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=NAVY)
TICKER_HDR = ("Ticker", "TKR", "Symbol")
SKIP_COLS = {"FMP financial read", "Key numbers (FMP)", "Strength %ile", "Name", "Company", "Ticker", "#", "Rank"}


def header_of(ws, max_row=12):
    """(header_row, {header: col}) for the first row with a Ticker cell."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row)):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in TICKER_HDR:
                return c.row, {str(x.value).strip(): x.column for x in row if x.value is not None}
    return None, {}


def insert_col(ws, idx, header, hdr_row, width=24):
    """Insert a column at idx, keeping merged ranges and column widths aligned
    (openpyxl's insert_cols moves cells but not merges or widths)."""
    widths = {i: ws.column_dimensions[get_column_letter(i)].width for i in range(1, ws.max_column + 1)}
    merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    ws.insert_cols(idx)
    for r0, c0, r1, c1 in merges:
        if c0 >= idx:
            c0, c1 = c0 + 1, c1 + 1
        elif c1 >= idx:
            c1 += 1
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
    for i in sorted(widths, reverse=True):
        j = i + 1 if i >= idx else i
        if widths[i]:
            ws.column_dimensions[get_column_letter(j)].width = widths[i]
    ws.column_dimensions[get_column_letter(idx)].width = width
    h = ws.cell(row=hdr_row, column=idx, value=header)
    h.font, h.fill = HEAD, FILL
    h.alignment = Alignment(wrap_text=True, vertical="center")


def _name_col(cols):
    for k in ("Name", "Company"):
        if k in cols:
            return cols[k]
    return cols.get("Ticker") or next(iter(cols.values()))


def _ticker(v):
    return v.replace("●", "").strip() if isinstance(v, str) else None


def key_line(f):
    if not f:
        return None
    bits = []
    if f.get("p_b") is not None:
        bits.append(f"P/B {f['p_b']:.2f}")
    ev = f.get("ev_ebitda")
    if (f.get("sector") or "") == "Financial Services":
        if f.get("pe") is not None and 0 < f["pe"] < 200:
            bits.append(f"P/E {f['pe']:.1f}")
    elif ev is not None and 0 < ev < 200:
        bits.append(f"{ev:.1f}× EBITDA")
    nc = f.get("net_cash_pct")
    if nc is not None and (f.get("sector") or "") != "Financial Services":
        bits.append(("cash " if nc >= 0 else "debt ") + f"{abs(nc) * 100:.0f}%")
    if f.get("range_pos") is not None:
        bits.append(f"52w {f['range_pos'] * 100:.0f}%")
    return " · ".join(bits)


def key_numbers(wb, fin, sym_map=None, skip=()):
    n = 0
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        hr, cols = header_of(ws)
        if not hr or "Key numbers (FMP)" in cols:
            continue
        idx = _name_col(cols) + 1
        insert_col(ws, idx, "Key numbers (FMP)", hr, width=30)
        tcol = cols.get("Ticker") or cols.get("TKR") or cols.get("Symbol")
        tcol = tcol + 1 if tcol >= idx else tcol
        for r in range(hr + 1, ws.max_row + 1):
            t = _ticker(ws.cell(row=r, column=tcol).value)
            if t:
                line = key_line(fin.get((sym_map or {}).get(t, t)))
                if line:
                    ws.cell(row=r, column=idx, value=line)
        n += 1
    return n


# ---------------------------------------------------------------- one scale
def _load(fn):
    p = ROOT / fn
    return json.loads(p.read_text()) if p.exists() else {}


def _field(fn, field, nested=None):
    d = _load(fn)
    if nested:
        d = d.get(nested) or {}
    out = {}
    for k, v in d.items():
        if isinstance(v, dict):
            x = v.get(field)
            if isinstance(x, (int, float)) and x > 0:
                out[k] = float(x)
    return out


def populations():
    """Tab -> {ticker: score} over EVERYTHING the layer scored."""
    return {
        "Governance Discount": _field("governance_discount.json", "score"),
        "Mechanism Gates": _field("mechanism_gates.json", "score"),
        "Payoff Geometry": _field("payoff_geometry.json", "score"),
        "MD&A Intent": _field("mda_scan.json", "score"),
        "Call Intent": _field("call_intent.json", "act_prob"),
        "Structured Distressed": _field("structured_distressed_injection.json", "score"),
        "Distressed Stub Progress": _field("distressed_stub_progress.json", "score"),
        "Re-Rate Catalysts": _field("rerate_catalysts.json", "rerate_score") or _field("rerate_catalysts.json", "catalyst_score"),
        "Tail Odds": _field("tail_odds.json", "est_tail_prob"),
        "Insider Conviction": _field("discretionary_insider_conviction.json", "score"),
        "Insider Filing-Time": _field("form4_timing.json", "score", nested="scores"),
        "Asymmetry Assembly": _field("asymmetry_assembly.json", "score"),
        "Foreign Markets": _field("foreign_markets.json", "score"),
    }


SCORE_HDRS = ("Score", "Conviction", "Timing pts", "Improve", "Assembly", "Conf. tail p", "Tail p",
              "Act prob", "Consensus", "Screens fired", "Distress")


def _pctile(sorted_vals, v):
    if not sorted_vals:
        return None
    lo, hi = bisect_left(sorted_vals, v), bisect_right(sorted_vals, v)
    return round(100.0 * ((lo + hi) / 2) / len(sorted_vals))


def strength(wb, pops, skip=()):
    """Insert 'Strength %ile' after the name (after Key numbers when present)."""
    n = 0
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        hr, cols = header_of(ws)
        if not hr or "Strength %ile" in cols:
            continue
        sc = next((cols[h] for h in SCORE_HDRS if h in cols), None)
        pop = pops.get(ws.title)
        if sc is None and not pop:
            continue
        tcol = cols.get("Ticker") or cols.get("TKR") or cols.get("Symbol")
        rows = [r for r in range(hr + 1, ws.max_row + 1) if _ticker(ws.cell(row=r, column=tcol).value)]
        if pop:
            vals, basis = sorted(pop.values()), f"all {len(pop):,} names scored"
            get = lambda r: pop.get(_ticker(ws.cell(row=r, column=tcol).value))
        else:
            within = [ws.cell(row=r, column=sc).value for r in rows]
            vals = sorted(v for v in within if isinstance(v, (int, float)))
            basis = f"the {len(vals)} names on this tab"
            get = lambda r: ws.cell(row=r, column=sc).value if isinstance(ws.cell(row=r, column=sc).value, (int, float)) else None
        if not vals:
            continue
        idx = (cols.get("Key numbers (FMP)") or _name_col(cols)) + 1
        scores = {r: get(r) for r in rows}
        insert_col(ws, idx, "Strength %ile", hr, width=9)
        ws.cell(row=hr, column=idx).comment = None
        for r, v in scores.items():
            if v is not None:
                ws.cell(row=r, column=idx, value=_pctile(vals, v))
        # say what the percentile is measured against, under the header row
        from openpyxl.cell.cell import MergedCell
        note = ws.cell(row=hr - 1, column=idx) if hr > 1 else None
        if note is not None and not isinstance(note, MergedCell) and note.value is None:
            note.value = f"%ile vs {basis}"
            note.font = Font(italic=True, size=8, color="666666")
        n += 1
    return n


# ---------------------------------------------------------------- grouping
GROUP_COLORS = {"Decide": "1F3864", "Theses": "2E75B6", "Signals": "70AD47",
                "Evidence": "ED7D31", "Plumbing": "A5A5A5"}


def regroup(wb, groups, contents="Contents", descriptions=None):
    """Reorder sheets by group; colour tabs; rewrite the Contents table."""
    order, seen = [], set()
    for g, tabs in groups:
        for t in tabs:
            if t in wb.sheetnames and t not in seen:
                order.append((g, t)); seen.add(t)
    for t in wb.sheetnames:                       # anything unlisted goes to Plumbing
        if t not in seen:
            order.append(("Plumbing", t)); seen.add(t)
    wb._sheets = [wb[t] for _, t in order]
    for g, t in order:
        wb[t].sheet_properties.tabColor = GROUP_COLORS.get(g, "A5A5A5")
    if contents in wb.sheetnames:
        ws = wb[contents]
        hr, cols = None, {}
        for row in ws.iter_rows(min_row=1, max_row=8):
            for c in row:
                if c.value == "Tab":
                    hr = c.row
            if hr:
                break
        if hr:
            old = {}
            for r in range(hr + 1, ws.max_row + 1):
                vals = [ws.cell(row=r, column=j).value for j in range(1, 5)]
                tab = next((v for v in vals if isinstance(v, str) and v in wb.sheetnames), None)
                if tab:
                    old[tab] = next((v for v in vals[vals.index(tab) + 1:] if isinstance(v, str)), "")
            for m in list(ws.merged_cells.ranges):
                if m.min_row >= hr:
                    ws.unmerge_cells(str(m))
            for r in range(hr, ws.max_row + 1):
                for j in range(1, 5):
                    ws.cell(row=r, column=j).value = None
            for j, h in enumerate(["Group", "Tab", "What it contains"], 1):
                c = ws.cell(row=hr, column=j, value=h)
                c.font, c.fill = HEAD, FILL
            ws.column_dimensions["A"].width = 11
            ws.column_dimensions["B"].width = 26
            ws.column_dimensions["C"].width = 90
            r = hr
            for g, t in order:
                if t == contents:
                    continue
                r += 1
                ws.cell(row=r, column=1, value=g).font = Font(bold=True, color=GROUP_COLORS.get(g, "000000"))
                ws.cell(row=r, column=2, value=t)
                ws.cell(row=r, column=3, value=(descriptions or {}).get(t) or old.get(t, ""))
    return order


# ---------------------------------------------------------------- tear sheets
METRICS = [("P/B", "p_b", "{:.2f}"), ("P/E", "pe", "{:.1f}"), ("EV/EBITDA", "ev_ebitda", "{:.1f}×"),
           ("FCF yield", "fcf_yield", "{:.0%}"), ("Net cash / mcap", "net_cash_pct", "{:+.0%}"),
           ("Net debt / EBITDA", "nd_ebitda", "{:.1f}×"), ("Revenue growth", "rev_growth", "{:+.0%}"),
           ("Operating margin", "op_m", "{:.0%}"), ("ROE", "roe", "{:.0%}"),
           ("Shares YoY", "shares_yoy", "{:+.0%}"), ("Interest cover", "int_cover", "{:.1f}×"),
           ("52-week position", "range_pos", "{:.0%}"), ("Avg $ volume", "adv_usd", "${:,.0f}")]


def _row_facts(ws, r, cols):
    out = []
    for h, c in cols.items():
        if h in SKIP_COLS or h in TICKER_HDR:
            continue
        v = ws.cell(row=r, column=c).value
        if v in (None, "", "–", "—"):
            continue
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            s = f"{v:,.0f}" if abs(v) >= 1000 else (f"{v:.3g}" if isinstance(v, float) else str(v))
        else:
            s = str(v)
        out.append(f"{h}: {s[:60]}")
    return " · ".join(out)


def tear_sheets(wb, fin, names, sym_map=None, title="Tear Sheets", index=None, max_names=60):
    """names: [display_ticker] in priority order."""
    # index every tab row by ticker once
    where = {}
    for ws in wb.worksheets:
        if ws.title in (title, "Name Financials"):
            continue
        hr, cols = header_of(ws)
        if not hr:
            continue
        tcol = cols.get("Ticker") or cols.get("TKR") or cols.get("Symbol")
        for r in range(hr + 1, ws.max_row + 1):
            t = _ticker(ws.cell(row=r, column=tcol).value)
            if t and len(t) < 20:
                where.setdefault(t, []).append((ws, r, cols))
    ts = wb.create_sheet(title, index) if index is not None else wb.create_sheet(title)
    ts.sheet_view.showGridLines = False
    ts.column_dimensions["A"].width = 24
    ts.column_dimensions["B"].width = 12
    ts.column_dimensions["C"].width = 110
    ts["A1"] = title
    ts["A1"].font = Font(bold=True, size=14, color=NAVY)
    ts["A2"] = ("One block per name: identity, the FMP financial panel, and every tab the name "
                "appears on -- with its strength percentile (100 = strongest in that layer's full "
                "population) and the facts that tab records. Names in priority order.")
    ts["A2"].alignment = Alignment(wrap_text=True)
    ts.merge_cells("A2:C2")
    ts.row_dimensions[2].height = 30
    r = 4
    done = 0
    for t in names:
        if done >= max_names or t not in where:
            continue
        sym = (sym_map or {}).get(t, t)
        f = fin.get(sym) or {}
        done += 1
        nm = f.get("name")
        if not nm:                                # no FMP record: take the name from the tab itself
            ws0, r0, c0 = where[t][0]
            nc = c0.get("Name") or c0.get("Company")
            nm = ws0.cell(row=r0, column=nc).value if nc else ""
        head = f"{t}  —  {nm or ''}"
        meta = " · ".join(x for x in [f.get("sector"), f.get("country"),
                                      f"mcap ${(f.get('mcap') or 0) / 1e6:,.0f}M" if f.get("mcap") else None,
                                      f"price {f.get('price')}" if f.get("price") else None] if x)
        c = ts.cell(row=r, column=1, value=head)
        c.font, c.fill = Font(bold=True, color="FFFFFF", size=12), FILL
        ts.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
        ts.cell(row=r, column=1, value=meta).font = Font(italic=True, color="444444")
        ts.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
        if f.get("read"):
            ts.cell(row=r, column=1, value="Financial read").font = Font(bold=True)
            ts.cell(row=r, column=2, value=f.get("read") + (
                "  [" + "; ".join(f["flags"]) + "]" if f.get("flags") else ""))
            ts.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            r += 1
        # metric grid: two metrics per row
        cells = [(lab, fmt.format(f[k])) for lab, k, fmt in METRICS
                 if f.get(k) is not None and not (k == "int_cover" and f[k] == 0)]
        for i in range(0, len(cells), 2):
            pair = cells[i:i + 2]
            ts.cell(row=r, column=1, value=pair[0][0]).font = Font(color="555555")
            ts.cell(row=r, column=2, value=pair[0][1])
            if len(pair) > 1:
                ts.cell(row=r, column=3, value=f"{pair[1][0]}: {pair[1][1]}")
            r += 1
        ts.cell(row=r, column=1, value="Appears on").font = Font(bold=True)
        ts.cell(row=r, column=2, value="Strength %ile").font = Font(bold=True)
        ts.cell(row=r, column=3, value="What that tab says").font = Font(bold=True)
        r += 1
        for ws, row, cols in where[t]:
            ts.cell(row=r, column=1, value=ws.title)
            st = cols.get("Strength %ile")
            ts.cell(row=r, column=2, value=ws.cell(row=row, column=st).value if st else None)
            c = ts.cell(row=r, column=3, value=_row_facts(ws, row, cols))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ts.row_dimensions[r].height = 30
            r += 1
        r += 1
    ts.freeze_panes = "A4"
    return done
