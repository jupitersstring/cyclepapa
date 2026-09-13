"""Truly-XR book — forensic CONFLUENCE ranking.

The Part II thesis operationalised: the grossest, least-arbitraged re-ratings
are not one gap but SEVERAL independent GAAP-vs-economic gaps STACKED in one
cheap name, converging on a MECHANICAL (self-executing) catalyst. This book
ranks names that carry >= 3 independent forensic tells (truly_xr_flag), by
truly_xr_score (tells x mechanical-catalyst upweight x cheapness), and shows the
SOURCE gaps per name (mechanical catalysts marked *) so the reader sees why it
ranks.

Output: truly_xr_book.xlsx
"""
from __future__ import annotations
import argparse
import os
import sys

import numpy as np
import pandas as pd

import tab_colors
from build_harvard_workbook import (
    Workbook,
    INK, MUTED, RULE,
    _font, _section_rule, _verdict_badge,
    _write_money, _write_pct, _write_score, _write_int,
    _TXT_ALIGN_LEFT, _NUM_ALIGN_CENTER,
)
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter


TRULY_COLS = ['symbol', 'truly_xr_score', 'truly_xr_flag', 'truly_xr_tell_count',
              'truly_xr_mech_count', 'truly_xr_tells_str', 'forensic_xr_score']


def load_data():
    df = pd.read_csv('asymmetry_global.csv', low_memory=False).drop_duplicates('symbol')
    if os.path.exists('archetype_tags.csv'):
        t = pd.read_csv('archetype_tags.csv', low_memory=False)
        keep = ['symbol'] + [c for c in TRULY_COLS if c != 'symbol' and c in t.columns]
        df = df.drop(columns=[c for c in keep if c != 'symbol' and c in df.columns],
                     errors='ignore')
        df = df.merge(t[keep].drop_duplicates('symbol'), on='symbol', how='left')

    # fresh verdicts (rolling diligence log wins over the enrich snapshot)
    frames = []
    for path, default in [('qualitative_aligned_green.csv', 'GREEN'),
                          ('qualitative_red_avoid.csv', 'RED'),
                          ('qualitative_extended_verdicts.csv', None)]:
        if not os.path.exists(path):
            continue
        try:
            d = pd.read_csv(path)
        except pd.errors.ParserError:
            d = pd.read_csv(path, engine='python', on_bad_lines='skip', quoting=3)
        if 'verdict' not in d.columns and default is not None:
            d['verdict'] = default
        frames.append(d[[c for c in ['symbol', 'verdict'] if c in d.columns]])
    if frames:
        v = pd.concat(frames, ignore_index=True).drop_duplicates('symbol', keep='last')
        df = df.drop(columns=[c for c in ('verdict',) if c in df.columns])
        df = df.merge(v, on='symbol', how='left')
    if 'verdict' not in df.columns:
        df['verdict'] = pd.NA
    df['verdict'] = df['verdict'].fillna('UNRESEARCHED')

    from otc_flag import apply_otc_mode as _apply_otc
    df = _apply_otc(df, 'ex-otc')
    if 'market_cap_usd' in df.columns:
        df['market_cap'] = (pd.to_numeric(df['market_cap_usd'], errors='coerce')
                            .fillna(pd.to_numeric(df.get('market_cap'), errors='coerce')))
    df = df[df['verdict'] != 'RED']
    df['truly_xr_score'] = pd.to_numeric(df.get('truly_xr_score'), errors='coerce').fillna(0)
    df['truly_xr_flag'] = pd.to_numeric(df.get('truly_xr_flag'), errors='coerce').fillna(0)
    return df[df['truly_xr_flag'] == 1].copy()


HEADERS = ['#', 'Ticker', 'Name', 'Country', 'Sector', 'Mcap (USD)', 'Verdict',
           'TrulyXR', 'Tells', 'Mech', 'ForXR', 'EV/EBITDA', 'P/B', 'FCF yld %',
           'Source gaps (mechanical *)']
NCOLS = len(HEADERS)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--n', type=int, default=250, help='top N by truly_xr_score')
    ap.add_argument('--out', default='truly_xr_book.xlsx')
    args = ap.parse_args()

    print('loading data...', file=sys.stderr)
    df = load_data()
    print(f'  {len(df):,} names flagged truly_xr (>=3 independent tells, cheap)', file=sys.stderr)
    df = df.sort_values('truly_xr_score', ascending=False).head(args.n)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Truly-XR'
    tab_colors.set_tab(ws, tab_colors.FAMILY_COLORS['deep_value'])

    f_bold = _font(bold=True, color=INK)
    f_italic_muted = _font(italic=True, color=MUTED)
    f_text = _font(color=INK)

    t = ws.cell(row=2, column=1, value='Truly-XR — forensic confluence (>= 3 independent tells, cheap)')
    t.font = f_bold
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOLS)
    ws.row_dimensions[2].height = 22
    sub = ws.cell(row=3, column=1,
                  value='Several independent GAAP-vs-economic gaps STACKED in one cheap, non-melting name; '
                        'ranked by tells x mechanical-catalyst upweight x cheapness. Source gaps shown per name (mechanical *).')
    sub.font = f_italic_muted
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NCOLS)
    _section_rule(ws, 5, 'Ranked by truly_xr_score — the confluence is the edge (Part II synthesis)',
                  span_cols=NCOLS)

    hdr = 6
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=hdr, column=i, value=h)
        c.font = _font(bold=True, color=MUTED)
        c.alignment = _TXT_ALIGN_LEFT if i in (2, 3, 4, 5, NCOLS) else _NUM_ALIGN_CENTER
    for c in range(1, NCOLS + 1):
        ws.cell(row=hdr + 1, column=c).border = Border(top=Side(style='thin', color=INK))

    r = hdr + 1
    for rank, (_, row) in enumerate(df.iterrows(), start=1):
        r += 1
        ws.cell(row=r, column=1, value=rank).font = f_text
        ws.cell(row=r, column=2, value=str(row.get('symbol') or '')).font = f_text
        ws.cell(row=r, column=3, value=str(row.get('name') or '')[:30]).font = f_text
        ws.cell(row=r, column=3).alignment = _TXT_ALIGN_LEFT
        ws.cell(row=r, column=4, value=str(row.get('src') or '')).font = _font(color=MUTED)
        ws.cell(row=r, column=5, value=str(row.get('sector') or '')[:14]).font = _font(color=MUTED)
        ws.cell(row=r, column=5).alignment = _TXT_ALIGN_LEFT
        _write_money(ws, r, 6, row.get('market_cap'), font=f_text)
        _v = row.get('verdict')
        _verdict_badge(ws, r, 7, _v if (pd.notna(_v) and _v) else 'UNRESEARCHED')
        _write_score(ws, r, 8, row.get('truly_xr_score'), font=f_text)
        _write_int(ws, r, 9, row.get('truly_xr_tell_count'), font=f_text)
        _write_int(ws, r, 10, row.get('truly_xr_mech_count'), font=f_text)
        _write_score(ws, r, 11, row.get('forensic_xr_score'), font=f_text)
        _write_score(ws, r, 12, row.get('ev_ebitda'), font=f_text)
        _write_score(ws, r, 13, row.get('pb'), font=f_text)
        _write_pct(ws, r, 14, row.get('fcf_yield'), font=f_text)
        _gaps = ws.cell(row=r, column=NCOLS, value=str(row.get('truly_xr_tells_str') or ''))
        _gaps.font = _font(color=INK)
        _gaps.alignment = _TXT_ALIGN_LEFT
        for c in range(1, NCOLS + 1):
            ws.cell(row=r, column=c).border = Border(bottom=Side(style='thin', color=RULE))

    widths = {1: 4, 2: 9, 3: 30, 4: 6, 5: 14, 6: 13, 7: 12, 8: 9, 9: 6, 10: 6,
              11: 9, 12: 10, 13: 7, 14: 10, 15: 46}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'A{hdr + 2}'
    if ws.max_row > hdr + 1:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(NCOLS)}{ws.max_row}"

    wb.save(args.out)
    print(f'wrote {args.out}: {len(df)} rows', file=sys.stderr)


if __name__ == '__main__':
    main()
