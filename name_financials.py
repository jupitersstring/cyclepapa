"""Per-name financial profile from FMP -- the numbers a reader needs next to
every name in the books.

For every symbol in the US quote store, the risk-reward book's mapped
symbols and the OTC intent names, one record from the FMP bulk files
(no per-name API calls):

  valuation   price, mcap, EV, P/B (validated: fmp_book), P/E, P/S, EV/EBITDA,
              FCF yield, earnings yield, dividend yield
  quality     gross / operating / net margin, ROE, ROIC, SBC % revenue
  balance     net cash % mcap (same-currency), net debt / EBITDA, debt/equity,
              interest cover, current ratio
  trajectory  revenue growth (TTM vs prior TTM) and share-count change (YoY)
              from the bulk quarterly statements
  trading     52-week range position, average $ volume
  read        a one-line plain-English summary + data-quality flags

Output: name_financials.json {symbol: {...}}.  Also provides
financials_sheet(wb, title, symbols, fin) used by all three workbooks.
"""

from __future__ import annotations

import csv
import glob
import json
import re
from pathlib import Path

import fmp_book
import fmp_client as fmp

ROOT = Path("/home/user/cyclepapa")
OUT = ROOT / "name_financials.json"


def _f(x):
    try:
        v = float(x)
        return v if v == v else None
    except (TypeError, ValueError):
        return None


def _quarterly(prefix, fields):
    """{sym: [(date, {field: value})...] ascending} from cached bulk statements."""
    out = {}
    for fn in glob.glob(str(ROOT / "fmp_cache" / f"{prefix}_*")):
        for r in csv.DictReader(open(fn, encoding="utf-8", errors="ignore")):
            if r.get("date"):
                out.setdefault(r["symbol"], {})[r["date"][:10]] = {k: _f(r.get(k)) for k in fields}
    return {s: sorted(m.items()) for s, m in out.items()}


def _days(a, b):
    from datetime import date
    return (date.fromisoformat(b) - date.fromisoformat(a)).days


def _ttm_growth(rows, key):
    """TTM vs prior TTM using statement DATES: quarterly reporters sum the last
    4 vs prior 4 quarters; semi-annual reporters (gaps ~6 months) the last 2
    vs prior 2 halves. Mixed / gappy histories return None."""
    pts = [(d, v[key]) for d, v in rows if v.get(key) is not None]
    if len(pts) < 4:
        return None
    gaps = [_days(pts[i][0], pts[i + 1][0]) for i in range(len(pts) - 1)]
    last = gaps[-3:]
    if all(70 <= g <= 110 for g in gaps[-7:]) and len(pts) >= 8:
        n = 4
    elif all(160 <= g <= 200 for g in last[-3:]) and len(pts) >= 4:
        n = 2
    else:
        return None
    vals = [v for _, v in pts]
    now, prev = sum(vals[-n:]), sum(vals[-2 * n:-n])
    return (now / prev - 1) if prev and prev > 0 else None


def _yoy(rows, key):
    """Value now vs the statement closest to one year earlier (330-400 days)."""
    pts = [(d, v[key]) for d, v in rows if v.get(key)]
    if len(pts) < 2:
        return None
    d_now, v_now = pts[-1]
    prior = [(abs(_days(d, d_now) - 365), v) for d, v in pts[:-1] if 330 <= _days(d, d_now) <= 400]
    if not prior:
        return None
    v0 = min(prior)[1]
    return (v_now / v0 - 1) if v0 else None


def _range_pos(rng, price):
    try:
        lo, hi = [float(x) for x in str(rng).split("-")[:2]]
        return (price - lo) / (hi - lo) if hi > lo and price else None
    except (ValueError, TypeError):
        return None


def universe():
    syms = set(json.loads((ROOT / "yfinance_quick.json").read_text()))
    for fn in ("otc_intent.json", "call_intent.json", "fmp_quotes.json", "governance_discount.json",
               "mechanism_gates.json", "rerate_catalysts.json", "event_detail.json", "psu_detail.json",
               "tail_odds.json", "payoff_geometry.json", "distressed_stub_progress.json", "foreign_markets.json"):
        p = ROOT / fn
        if p.exists():
            syms |= set(json.loads(p.read_text()))
    return syms


_BULK = None


def _bulk():
    global _BULK
    if _BULK is None:
        prof = {}
        for fn in glob.glob(str(ROOT / "fmp_cache" / "profile-bulk_part*.csv")):
            for r in csv.DictReader(open(fn, encoding="utf-8", errors="ignore")):
                prof[r["symbol"]] = r
        _BULK = (prof, {r["symbol"]: r for r in fmp.get_bulk_csv("ratios-ttm-bulk")},
                 {r["symbol"]: r for r in fmp.get_bulk_csv("key-metrics-ttm-bulk")},
                 fmp_book.load(), _quarterly("isbulk", ["revenue", "weightedAverageShsOutDil", "ebitda", "netInterestIncome",
                                                        "netIncome", "depreciationAndAmortization",
                                                        "netIncomeFromDiscontinuedOperations"]),
                 _quarterly("cfbulk", ["commonStockRepurchased", "commonDividendsPaid", "commonStockIssuance"]))
    return _BULK


def security_type(sym, p, prof):
    """common / note-preferred / fund / spac / defunct -- only common equity belongs on an equity tab."""
    nm = p.get("companyName") or ""
    ind = p.get("industry") or ""
    if re.search(r"\d%|\bnotes?\b|debenture|\bpfd\b|preferred|\bZONES\b|capital securities|cap secs", nm, re.I):
        return "note / preferred line"
    if str(p.get("isFund")).lower() == "true" or str(p.get("isEtf")).lower() == "true" or (
            ind.startswith("Asset Management") and re.search(r"\bFund\b|\bTrust\b|Income|Municipal|Opportunit|Strategy|Portfolio", nm)
            and not re.search(r"\bInc\.?$|Corp|Group|Holdings|Management", nm)):
        return "fund"
    if ind == "Shell Companies" and re.search(r"Acquisition|Capital Corp|SPAC|Merger|Blank Check", nm, re.I):
        return "spac"
    price = _f(p.get("price"))
    q = prof.get(sym + "Q")
    same_issuer = bool(q) and str(q.get("isActivelyTrading")).lower() != "false" and (_f(q.get("price")) or 0) < 5 and ((p.get("cik") and p.get("cik") == q.get("cik")) or
                               re.sub(r"[^a-z]", "", (q.get("companyName") or "").lower())[:20]
                               == re.sub(r"[^a-z]", "", nm.lower())[:20])
    if same_issuer or (str(p.get("isActivelyTrading")).lower() == "false" and (price or 0) < 1):
        return "bankrupt / not trading"          # only when the Q line is the SAME issuer (BACQ is not Bank of America)
    return "common"


def kind_of(sector, industry):
    i = industry or ""
    if i.startswith("Banks") or i in ("Financial - Mortgages", "Financial - Credit Services"):
        return "bank"
    if i.startswith("Insurance") and "Brokers" not in i:
        return "insurer"
    if i == "REIT - Mortgage":
        return "mreit"
    if i.startswith("REIT"):
        return "reit"
    if i.startswith("Asset Management") or i in ("Financial - Capital Markets", "Investment - Banking & Investment Services",
                                                 "Financial - Conglomerates", "Shell Companies"):
        return "financial_other"
    if i in ("Oil & Gas Exploration & Production", "Oil & Gas Integrated"):
        return "ep"
    return "operating"


def _ttm_sum(qs, key):
    """Trailing-12-month sum from dated quarterly rows (semi-annual filers: 2 halves)."""
    if not qs:
        return None
    last = qs[-1][0]
    rows = [(d, v.get(key)) for d, v in qs if _days(d, last) < 360 and v.get(key) is not None]
    if len(rows) >= 4:
        return sum(v for _, v in rows[-4:])
    if len(rows) == 2 and _days(rows[0][0], rows[1][0]) > 150:
        return sum(v for _, v in rows)
    return None


def _usd(v, ccy):
    """Local-currency amount in USD (minor units GBp / ZAc / ILA handled)."""
    if v is None or not ccy:
        return v if ccy in (None, "", "USD") else None
    minor = {"GBp": ("GBP", 1), "GBX": ("GBP", 1), "ZAc": ("ZAR", 1), "ZAC": ("ZAR", 1), "ILA": ("ILS", 1)}
    c = minor.get(ccy, (ccy, 1))[0]                   # FMP marketCap is in MAJOR units already
    if c == "USD":
        return v
    r = fmp_book.fx(c, "USD")
    return v * r if r else None


def build(symbols) -> dict:
    """Financial records for the given FMP symbols (those FMP knows)."""
    prof, rat, km, sheets, inc, cfq = _bulk()
    out = {}
    for s in sorted(set(symbols) & set(prof)):
        p, ra, k = prof[s], rat.get(s) or {}, km.get(s) or {}
        price, mcap = _f(p.get("price")), _f(p.get("marketCap"))
        pb, pb_src = fmp_book.pb(s, mcap, p.get("currency"), ra.get("priceToBookRatioTTM"), sheets,
                                 price, str(p.get("isAdr")).lower() == "true")
        bs = sheets.get(s) or {}
        eq = _f(bs.get("totalStockholdersEquity"))
        cash, debt = _f(bs.get("cashAndShortTermInvestments")), _f(bs.get("totalDebt"))
        # net cash as % of mcap inside ONE statement currency: (cash-debt)/equity x 1/(P/B)
        ncp = ((cash or 0) - (debt or 0)) / eq / pb if (eq and eq > 0 and pb and cash is not None) else None
        q = inc.get(s) or []
        sh_chg = _yoy(q, "weightedAverageShsOutDil")
        if sh_chg is not None and not (-0.6 < sh_chg < 3):
            sh_chg = None                                   # split artefact
        rec = {
            "name": p.get("companyName"), "sector": p.get("sector"), "country": p.get("country"),
            "currency": p.get("currency"), "price": price, "mcap": mcap, "mcap_usd": _usd(mcap, p.get("currency")),
            "ev": _f(k.get("enterpriseValueTTM")), "p_b": pb, "pb_src": pb_src,
            "pe": _f(ra.get("priceToEarningsRatioTTM")), "ps": _f(ra.get("priceToSalesRatioTTM")),
            "ev_ebitda": _f(k.get("evToEBITDATTM")), "fcf_yield": _f(k.get("freeCashFlowYieldTTM")),
            "earn_yield": _f(k.get("earningsYieldTTM")), "div_yield": _f(ra.get("dividendYieldTTM")),
            "gross_m": _f(ra.get("grossProfitMarginTTM")), "op_m": _f(ra.get("operatingProfitMarginTTM")),
            "net_m": _f(ra.get("netProfitMarginTTM")), "roe": _f(k.get("returnOnEquityTTM")),
            "roic": _f(k.get("returnOnInvestedCapitalTTM")),
            "sbc_rev": _f(k.get("stockBasedCompensationToRevenueTTM")),
            "net_cash_pct": ncp, "nd_ebitda": _f(k.get("netDebtToEBITDATTM")),
            "de": _f(ra.get("debtToEquityRatioTTM")), "int_cover": _f(ra.get("interestCoverageRatioTTM")),
            "current": _f(ra.get("currentRatioTTM")),
            "rev_growth": _ttm_growth(q, "revenue"), "shares_yoy": sh_chg,
            "range_pos": _range_pos(p.get("range"), price),
            "adv_usd": _usd((_f(p.get("averageVolume")) or 0) * (price or 0) / (100 if p.get("currency") in ("GBp", "GBX", "ZAc", "ZAC", "ILA") else 1),
                            p.get("currency")) or None,
            "stmt_date": bs.get("date"),
        }
        rec["industry"] = p.get("industry")
        rec["kind"] = kind_of(p.get("sector"), p.get("industry"))
        # ---- sector-routed metrics (statement currency, expressed vs equity / mcap)
        ta = _f(bs.get("totalAssets"))
        gw = _f(bs.get("goodwillAndIntangibleAssets")) or 0
        pref = _f(bs.get("preferredStock")) or 0
        tang = (eq - gw - pref) if eq else None
        mc_stmt = (eq * pb) if (eq and pb) else None            # market cap in statement currency
        ttm = lambda qs, k: _ttm_sum(qs, k)
        ni, nii, da = ttm(q, "netIncome"), ttm(q, "netInterestIncome"), ttm(q, "depreciationAndAmortization")
        rec["tce_ta"] = tang / ta if (tang is not None and ta) else None
        rec["p_tbv"] = mc_stmt / tang if (mc_stmt and tang and tang > 0) else None
        aoci = _f(bs.get("accumulatedOtherComprehensiveIncomeLoss"))
        rec["aoci_eq"] = aoci / eq if (aoci is not None and eq and eq > 0) else None
        rec["roa"] = ni / ta if (ni is not None and ta) else None
        rec["nii_assets"] = nii / ta if (nii and ta and rec["kind"] in ("bank",)) else None
        rec["ffo_yield"] = (ni + da) / mc_stmt if (rec["kind"] == "reit" and ni is not None and da is not None and mc_stmt) else None
        lti = _f(bs.get("longTermInvestments"))
        rec["lt_inv_mcap"] = lti / mc_stmt if (lti and mc_stmt and rec["kind"] == "operating") else None
        rec["tax_assets_mcap"] = (_f(bs.get("taxAssets")) or 0) / mc_stmt if mc_stmt else None
        cq = cfq.get(s) or []
        rep, divp = ttm(cq, "commonStockRepurchased"), ttm(cq, "commonDividendsPaid")
        rec["buyback_ttm_mcap"] = abs(rep) / mc_stmt if (rep is not None and mc_stmt) else None
        rec["div_paid_ttm_mcap"] = abs(divp) / mc_stmt if (divp is not None and mc_stmt) else None
        for k_ in ("buyback_ttm_mcap", "div_paid_ttm_mcap"):
            if rec[k_] is not None and rec[k_] > 0.5:
                rec[k_] = None                            # > 50% of mcap in a year: statement / mcap basis mismatch
        if rec["kind"] in ("bank", "insurer", "mreit"):
            rec["net_cash_pct"] = rec["nd_ebitda"] = rec["int_cover"] = rec["ev_ebitda"] = None   # a lender's debt is its raw material
        # sanity: FMP TTM ratios break on the same artefacts as P/B
        flags = []
        if pb_src in ("mcap_suspect", "implausible"):
            flags.append("market-cap/share data inconsistent")
            # every market-cap-based multiple inherits the bad market cap
            for k_ in ("mcap_usd", "pe", "ps", "ev_ebitda", "fcf_yield", "earn_yield", "net_cash_pct", "p_tbv",
                       "buyback_ttm_mcap", "div_paid_ttm_mcap", "lt_inv_mcap", "ffo_yield"):
                rec[k_] = None
        if pb_src == "neg_equity":
            flags.append("negative equity")
        if rec["pe"] is not None and rec["earn_yield"] is not None and (rec["pe"] > 0) != (rec["earn_yield"] > 0):
            flags.append("P/E vs earnings-yield sign mismatch")
        fin_co = (p.get("sector") or "") == "Financial Services"
        if fin_co:
            rec["fcf_yield"] = None                      # bank / insurer cash flow is loan and float flow, not FCF
        elif rec["fcf_yield"] is not None and abs(rec["fcf_yield"]) > 1.0:
            flags.append("FCF yield >100% (check)")
            rec["fcf_yield"] = None
        sec = security_type(s, p, prof)
        if sec != "common":
            flags.append(f"not common equity ({sec})")
            rec["not_common"] = True
        rec["security"] = sec
        if rec["roe"] is not None and abs(rec["roe"]) > 1.5:
            flags.append(f"ROE {rec['roe'] * 100:,.0f}% not meaningful (tiny or bad equity figure)")
            rec["roe"] = None
        if rec["pe"] is not None and 0 <= rec["pe"] < 1:
            rec["pe"] = None                             # sub-1x P/E is a units artefact
        # EV/EBITDA is meaningless when either side is negative (FMP reports -/- as a positive multiple)
        ebitda_ttm = _ttm_sum(q, "ebitda")
        if rec["ev_ebitda"] is not None and ((rec["ev"] is not None and rec["ev"] <= 0) or
                                             (ebitda_ttm is not None and ebitda_ttm <= 0) or rec["ev_ebitda"] <= 0):
            rec["ev_ebitda"] = None
        # P/B / P/E must imply the ROE the statements show; a 3x+ gap means the market cap is
        # on the wrong share basis (ADR ratio, stale count, preferred line priced as the company)
        pb_, pe_, roe_ = rec["p_b"], rec["pe"], rec["roe"]
        if pb_ and pe_ and 0 < pe_ < 200 and roe_ and roe_ > 0.02:
            gap = (pb_ / pe_) / roe_
            if gap < 0.33 or gap > 3.0:
                flags.append(f"valuation inputs inconsistent (P/B÷P/E implies ROE {pb_ / pe_ * 100:.0f}% vs {roe_ * 100:.0f}%: "
                             "ADR ratio / share-count basis)")
                rec["p_b"] = rec["pe"] = rec["mcap_usd"] = rec["p_tbv"] = None
                rec["pb_src"] = "inconsistent"
        if rec["mcap_usd"] is not None and rec["mcap_usd"] > 5e12:
            flags.append("market cap implausible (units)")
            rec["mcap_usd"] = None
        if rec["shares_yoy"] is not None and abs(rec["shares_yoy"]) > 0.15 and (p.get("country") or "US") != "US":
            flags.append("share-count history inconsistent (foreign line)")
            rec["shares_yoy"] = None
        if rec["shares_yoy"] is not None and rec["shares_yoy"] > 1.0:
            flags.append("new entity / merger: history not comparable")
            rec["shares_yoy"] = rec["rev_growth"] = None
        if rec["rev_growth"] is not None and rec["rev_growth"] <= -0.95:
            rec["rev_growth"] = None                      # stub period, not a collapse
        if rec["div_yield"] is not None and rec["div_yield"] > 0.25:
            flags.append("dividend yield >25% (special or bad data)")
            rec["div_yield"] = None
        if rec["rev_growth"] is not None and rec["rev_growth"] > 3.0:
            flags.append("revenue not comparable YoY (fair-value / one-off income?)")
            rec["rev_growth"] = None
        rec["flags"] = flags
        rec["read"] = read_line(rec)
        out[s] = rec
    return out


def main() -> int:
    import sys
    syms = universe() | set(sys.argv[1:])
    out = build(syms)
    OUT.write_text(json.dumps(out))
    print(f"wrote {OUT.name}: {len(out)} names")
    return 0


def _pct(x, d=0):
    return None if x is None else f"{x * 100:+.{d}f}%"


def read_line(r):
    """Plain-English one-liner a reader would write in the margin."""
    bits = []
    k = r.get("kind")
    bb = r.get("buyback_ttm_mcap")
    bb_txt = f"bought back {bb * 100:.1f}% of mcap (TTM)" if bb and bb >= 0.01 else None
    if k == "bank":
        if r.get("p_tbv"):
            bits.append(f"P/TBV {r['p_tbv']:.2f}")
        elif r.get("p_b"):
            bits.append(f"P/B {r['p_b']:.2f}")
        if r.get("pe") and 0 < r["pe"] < 200:
            bits.append(f"P/E {r['pe']:.1f}")
        if r.get("roa") is not None:
            bits.append(f"ROA {r['roa'] * 100:.2f}%")
        if r.get("tce_ta") is not None:
            bits.append(f"TCE/TA {r['tce_ta'] * 100:.1f}%" + (" ⚠" if r["tce_ta"] < 0.06 else ""))
        if r.get("nii_assets"):
            bits.append(f"NII/assets {r['nii_assets'] * 100:.2f}%")
        if r.get("aoci_eq") is not None and r["aoci_eq"] <= -0.05:
            bits.append(f"AOCI loss {abs(r['aoci_eq']) * 100:.0f}% of equity" + (" ⚠" if r["aoci_eq"] <= -0.15 else ""))
        if r.get("div_yield"):
            bits.append(f"div {r['div_yield'] * 100:.1f}%")
        if bb_txt:
            bits.append(bb_txt)
        return " · ".join(bits) + " (bank basis)"
    if k in ("insurer", "mreit", "financial_other"):
        if r.get("p_b") is not None:
            bits.append(f"P/B {r['p_b']:.2f}")
        if r.get("p_tbv") and k == "insurer" and r["p_b"] and r["p_tbv"] > r["p_b"] * 1.15:
            bits.append(f"P/TBV {r['p_tbv']:.2f}")
        if r.get("pe") and 0 < r["pe"] < 200:
            bits.append(f"P/E {r['pe']:.1f}")
        if r.get("roe") is not None:
            bits.append(f"ROE {r['roe'] * 100:.0f}%")
        if r.get("aoci_eq") is not None and r["aoci_eq"] <= -0.10 and k == "insurer":
            bits.append(f"AOCI loss {abs(r['aoci_eq']) * 100:.0f}% of equity")
        if r.get("div_yield"):
            bits.append(f"div {r['div_yield'] * 100:.1f}%")
        if bb_txt:
            bits.append(bb_txt)
        return " · ".join(bits) + {"insurer": " (insurer: book/earnings basis)", "mreit": " (mortgage REIT: book/dividend basis)",
                                   "financial_other": " (financial: book/earnings basis)"}[k]
    if k == "reit":
        if r.get("p_b") is not None:
            bits.append(f"P/B {r['p_b']:.2f}")
        if r.get("ffo_yield"):
            bits.append(f"FFO≈ yield {r['ffo_yield'] * 100:.1f}%")
        if r.get("div_yield"):
            bits.append(f"div {r['div_yield'] * 100:.1f}%")
        if r.get("nd_ebitda") is not None and 0 < r["nd_ebitda"] < 40:
            bits.append(f"ND/EBITDA {r['nd_ebitda']:.1f}×")
        if r.get("int_cover") not in (None, 0) and r["int_cover"] < 1.5:
            bits.append(f"interest cover {r['int_cover']:.1f}× ⚠")
        if bb_txt:
            bits.append(bb_txt)
        return " · ".join(bits) + " (REIT basis)"
    if (r.get("sector") or "") == "Financial Services":
        # banks / insurers / asset managers: cash-flow and EBITDA metrics are meaningless
        if r["p_b"] is not None:
            bits.append(f"P/B {r['p_b']:.2f}")
        if r["pe"] is not None and 0 < r["pe"] < 200:
            bits.append(f"P/E {r['pe']:.1f}")
        if r["roe"] is not None:
            bits.append(f"ROE {r['roe'] * 100:.0f}%")
        if r["div_yield"]:
            bits.append(f"div yld {r['div_yield'] * 100:.1f}%")
        if r["shares_yoy"] is not None and abs(r["shares_yoy"]) >= 0.02:
            bits.append(("buying back " if r["shares_yoy"] < 0 else "diluting ")
                        + f"{abs(r['shares_yoy']) * 100:.0f}%/yr")
        return " · ".join(bits) + " (financial: book/earnings basis)"
    if r["p_b"] is not None:
        bits.append(f"P/B {r['p_b']:.2f}")
    if r["ev_ebitda"] is not None and 0 < r["ev_ebitda"] < 200:
        bits.append(f"EV/EBITDA {r['ev_ebitda']:.1f}×")
    elif r["ps"] is not None:
        bits.append(f"P/S {r['ps']:.1f}")
    if r["fcf_yield"] is not None:
        bits.append(f"FCF yld {r['fcf_yield'] * 100:.0f}%")
    if r["net_cash_pct"] is not None:
        bits.append(("net cash " if r["net_cash_pct"] >= 0 else "net debt ")
                    + f"{abs(r['net_cash_pct']) * 100:.0f}% of mcap")
    if r["rev_growth"] is not None:
        bits.append(f"rev {r['rev_growth'] * 100:+.0f}%")
    if r["op_m"] is not None and abs(r["op_m"]) < 5:
        bits.append(f"op margin {r['op_m'] * 100:.0f}%")
    if r["shares_yoy"] is not None and abs(r["shares_yoy"]) >= 0.02:
        bits.append(("buying back " if r["shares_yoy"] < 0 else "diluting ")
                    + f"{abs(r['shares_yoy']) * 100:.0f}%/yr")
    if r["int_cover"] not in (None, 0) and r["int_cover"] < 1.5 and (r["net_cash_pct"] or 0) < -0.10:
        bits.append(f"interest cover {r['int_cover']:.1f}× ⚠")
    if bb_txt:
        bits.append(bb_txt)
    if r.get("lt_inv_mcap") and r["lt_inv_mcap"] >= 0.3:
        bits.append(f"LT investments {r['lt_inv_mcap'] * 100:.0f}% of mcap")
    return " · ".join(bits)


# ---------------------------------------------------------------- shared sheet
COLS = [("Ticker", 9), ("Name", 22), ("Sector", 14), ("Mcap $M", 9), ("P/B", 6), ("P/E", 6),
        ("EV/EBITDA", 8), ("FCF yld", 7), ("Net cash/mcap", 9), ("ND/EBITDA", 8),
        ("Rev g", 7), ("Op m", 7), ("ROE", 7), ("Shares YoY", 8), ("Int cover", 7),
        ("52w pos", 7), ("ADV $k", 8), ("Read", 70), ("Data flags", 26), ("Appears on", 60)]


def financials_sheet(wb, title, symbols, fin, subtitle="", index=None, tabs=None):
    """Add a financials sheet for `symbols` [(display_ticker, fmp_symbol)]."""
    import book_layout as bl
    k = bl.kit(wb)
    ws = wb.create_sheet(title, index) if index is not None else wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    k.title(ws.cell(row=1, column=1, value=title))
    k.subtitle(ws.cell(row=2, column=1, value=subtitle or (
        "FMP TTM ratios + bulk statements. P/B validated (market cap ÷ latest equity, FX-converted); "
        "net cash measured inside one statement currency; revenue growth = TTM vs prior TTM; shares "
        "YoY from diluted weighted shares.")))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
    ws.row_dimensions[2].height = 30
    for j, (h, w) in enumerate(COLS, 1):
        c = ws.cell(row=4, column=j, value=h)
        k.header(c)
        ws.column_dimensions[c.column_letter].width = w
    if k.header_height:
        ws.row_dimensions[4].height = k.header_height

    def num(x, d=1):
        return None if x is None else round(x, d)

    def pc(x):
        return None if x is None else f"{x * 100:.0f}%"
    r = 4
    seen = set()
    for disp, s in symbols:
        f = fin.get(s)
        if not f or s in seen:
            continue
        seen.add(s)
        r += 1
        vals = [disp, (f.get("name") or "")[:22], (f.get("sector") or "")[:14],
                num((f.get("mcap_usd") or 0) / 1e6, 0) if f.get("mcap_usd") else "—", num(f.get("p_b"), 2), num(f.get("pe"), 1),
                num(f.get("ev_ebitda"), 1), pc(f.get("fcf_yield")), pc(f.get("net_cash_pct")),
                num(f.get("nd_ebitda"), 1), pc(f.get("rev_growth")), pc(f.get("op_m")), pc(f.get("roe")),
                pc(f.get("shares_yoy")), num(f.get("int_cover"), 1), pc(f.get("range_pos")),
                num((f.get("adv_usd") or 0) / 1e3, 0), f.get("read"), "; ".join(f.get("flags") or []),
                ", ".join((tabs or {}).get(s, []))]
        band = (r - 4) % 2 == 0
        for j, v in enumerate(vals, 1):
            k.body(ws.cell(row=r, column=j, value=v), band=band, bold=(j == 1))
    ws.freeze_panes = "B5"
    return r - 4


def annotate_workbook(wb, fin, sym_map=None, skip=()):
    """Append an 'FMP financial read' column to every sheet whose header row
    has a 'Ticker' cell, so each listed name carries its numbers. sym_map maps
    a displayed ticker to its FMP symbol (default: identity)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    n = 0
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        hdr_row = tcol = None
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ("Ticker", "TKR", "Symbol"):
                    hdr_row, tcol = c.row, c.column
                    break
            if hdr_row:
                break
        if not hdr_row:
            continue
        col = ws.max_column + 1
        import book_layout as bl
        for rr_ in range(hdr_row, ws.max_row + 1):          # inherit the table's styling
            bl.clone(ws.cell(row=rr_, column=col - 1), ws.cell(row=rr_, column=col))
        h = ws.cell(row=hdr_row, column=col, value="FMP financial read")
        ws.column_dimensions[h.column_letter].width = 62
        hit = 0
        seen = getattr(wb, "_fin_seen", [])
        for r in range(hdr_row + 1, ws.max_row + 1):
            v = ws.cell(row=r, column=tcol).value
            if not isinstance(v, str) or not v.strip():
                continue
            key = v.replace("●", "").strip()
            sym = (sym_map or {}).get(key, key)
            f = fin.get(sym)
            if f:
                seen.append((key, sym))
                tabs = getattr(wb, "_fin_tabs", {})
                tabs.setdefault(sym, [])
                if ws.title not in tabs[sym]:
                    tabs[sym].append(ws.title)
                wb._fin_tabs = tabs
                ws.cell(row=r, column=col, value=f.get("read") + (
                    "  [" + "; ".join(f["flags"]) + "]" if f.get("flags") else ""))
                hit += 1
        wb._fin_seen = seen
        n += hit > 0
    return n


def add_financials(wb, fin, sym_map=None, skip=(), title="Name Financials", index=None):
    """Annotate every ticker sheet, then add one full financials sheet for
    every name that appears anywhere in the workbook."""
    n = annotate_workbook(wb, fin, sym_map, skip)
    rows = list(dict.fromkeys(getattr(wb, "_fin_seen", [])))
    tabs = getattr(wb, "_fin_tabs", {})
    # most-cited names first: the ones that appear on the most tabs
    rows.sort(key=lambda kv: -len(tabs.get(kv[1], [])))
    k = financials_sheet(wb, title, rows, fin, index=index, tabs=tabs,
                         subtitle="One row per name in this book, most-cited first. 'Appears on' lists every tab "
                                  "the name is on. FMP TTM ratios + bulk statements; P/B validated (market cap ÷ "
                                  "latest equity, FX-converted); net cash inside one statement currency.")
    print(f"  financials: read column on {n} sheets; {title} sheet with {k} names")
    return k


def load():
    return json.loads(OUT.read_text()) if OUT.exists() else {}


if __name__ == "__main__":
    raise SystemExit(main())
