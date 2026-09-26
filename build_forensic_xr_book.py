"""Forensic-XR asymmetry book.

Ranks the universe by forensic_xr_score — size-normalized hidden/unpriced
forensic value (LIFO reserve, pension surplus, reversible DTA valuation
allowance, off-BS equity-method stakes, off-EV hidden assets, RPO backlog
gross profit), each as a fraction of market cap, combined with valuation
(x1.5 when the visible business is also cheap) and gated to a monetizable,
tradeable, non-melting opportunity. Each row breaks down the SOURCE of the
hidden value so the reader sees why it ranks.

Output: forensic_xr_book.xlsx
"""
from __future__ import annotations
import argparse
import os
import sys

import numpy as np
import pandas as pd

import tab_colors
from build_harvard_workbook import (
    Workbook,
    INK, MUTED, RULE,
    _font, _section_rule, _verdict_badge,
    _write_money, _write_pct, _write_score, _write_int,
    _TXT_ALIGN_LEFT, _NUM_ALIGN_CENTER,
)
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter


def load_data():
    """asymmetry_global + the forensic score/columns from archetype_tags.csv."""
    df = pd.read_csv('asymmetry_global.csv', low_memory=False).drop_duplicates('symbol')
    # forensic_xr_score / forensic_hidden_pct live in the tags output
    if os.path.exists('archetype_tags.csv'):
        t = pd.read_csv('archetype_tags.csv', low_memory=False)
        # (audit P0-2) carry the data-quality flag so DQ-failed rows are barred
        _fcols = ('forensic_xr_score', 'forensic_hidden_pct', 'data_quality_flag') + ACCT_COLS
        keep = ['symbol'] + [c for c in _fcols if c in t.columns]
        df = df.drop(columns=[c for c in _fcols if c in df.columns], errors='ignore')
        df = df.merge(t[keep].drop_duplicates('symbol'), on='symbol', how='left')
        if 'data_quality_flag' in df.columns:
            df = df[pd.to_numeric(df['data_quality_flag'], errors='coerce').fillna(0) != 1]

    # fresh verdicts (rolling diligence log wins over the enrich snapshot)
    frames = []
    for path, default in [('qualitative_aligned_green.csv', 'GREEN'),
                          ('qualitative_red_avoid.csv', 'RED'),
                          ('qualitative_extended_verdicts.csv', None)]:
        if not os.path.exists(path):
            continue
        try:
            d = pd.read_csv(path)
        except pd.errors.ParserError:
            d = pd.read_csv(path, engine='python', on_bad_lines='skip', quoting=3)
        if 'verdict' not in d.columns and default is not None:
            d['verdict'] = default
        frames.append(d[[c for c in ['symbol', 'verdict'] if c in d.columns]])
    if frames:
        v = pd.concat(frames, ignore_index=True).drop_duplicates('symbol', keep='last')
        df = df.drop(columns=[c for c in ('verdict',) if c in df.columns])
        df = df.merge(v, on='symbol', how='left')
    if 'verdict' not in df.columns:
        df['verdict'] = pd.NA
    df['verdict'] = df['verdict'].fillna('UNRESEARCHED')

    from otc_flag import apply_otc_mode as _apply_otc
    df = _apply_otc(df, 'ex-otc')
    if 'market_cap_usd' in df.columns:
        df['market_cap'] = (pd.to_numeric(df['market_cap_usd'], errors='coerce')
                            .fillna(pd.to_numeric(df.get('market_cap'), errors='coerce')))
    df = df[df['verdict'] != 'RED']
    df['forensic_xr_score'] = pd.to_numeric(df.get('forensic_xr_score'), errors='coerce').fillna(0)
    return df[df['forensic_xr_score'] > 0].copy()


HEADERS = ['#', 'Ticker', 'Name', 'Country', 'Sector', 'Industry', 'Mcap (USD)', 'Verdict',
           'ForXR score', 'Hidden %mcap',
           'Pension %', 'DTA allow %', 'LIFO %', 'Eq-stake %', 'RPO %rev',
           'EV/EBITDA', 'P/B', 'FCF yld %',
           # accounting-quality block (FMP quarterly 3-statement forensics)
           'Beneish M', 'Accruals %TA', 'CFO/NI', 'Red', 'Green', 'Accounting check']
ACCT_COLS = ('fq_beneish_m', 'fq_sloan_accruals', 'fq_cfo_to_ni', 'fq_forensic_red_count',
             'fq_forensic_green_count', 'fq_forensic_clean_confirm',
             'fq_beneish_risk_flag', 'fq_beneish_ma_distorted', 'fq_high_accruals_flag',
             'fq_receivables_divergence_flag', 'fq_inventory_build_flag', 'fq_sbc_heavy_flag',
             'fq_cash_leads_earnings_flag', 'fq_wc_release_flag', 'fq_deleveraging_flag',
             'fq_gm_inflection_flag', 'fq_defrev_build_flag', 'fq_cash_tax_shield_flag',
             'fq_rev_yoy_streak_m', 'fq_netdebt_decline_months', 'fmp_q_status')
_TELLS = [('fq_beneish_risk_flag', 'Beneish>-1.78', 'red'), ('fq_beneish_ma_distorted', 'Beneish(M&A)', 'red'),
          ('fq_high_accruals_flag', 'high accruals', 'red'),
          ('fq_receivables_divergence_flag', 'receivables>sales', 'red'),
          ('fq_inventory_build_flag', 'inventory build', 'red'), ('fq_sbc_heavy_flag', 'SBC>=30% CFO', 'red'),
          ('fq_cash_leads_earnings_flag', 'FCF leads NI', 'green'), ('fq_wc_release_flag', 'WC release', 'green'),
          ('fq_deleveraging_flag', 'deleveraging', 'green'), ('fq_gm_inflection_flag', 'GM inflecting', 'green'),
          ('fq_defrev_build_flag', 'def-rev building', 'green'), ('fq_cash_tax_shield_flag', 'cash-tax shield', 'green')]


def acct_check(row) -> str:
    """Plain-language accounting read: red tells first, then green; 'clean'
    when the Beneish / accrual / divergence tests all pass; '' if unscored."""
    if str(row.get('fmp_q_status') or '') != 'ok':
        return ''
    def on(c):
        v = pd.to_numeric(row.get(c), errors='coerce')
        return bool(v == 1) if pd.notna(v) else False
    red = [lbl for c, lbl, k in _TELLS if k == 'red' and on(c)]
    grn = [lbl for c, lbl, k in _TELLS if k == 'green' and on(c)]
    parts = []
    if red:
        parts.append('WARN: ' + ', '.join(red))
    if grn:
        parts.append('+ ' + ', '.join(grn))
    if not parts and on('fq_forensic_clean_confirm'):
        parts.append('clean')
    return ' | '.join(parts)
NCOLS = len(HEADERS)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--n', type=int, default=250, help='top N by forensic_xr_score')
    ap.add_argument('--out', default='forensic_xr_book.xlsx')
    args = ap.parse_args()

    print('loading data...', file=sys.stderr)
    df = load_data()
    print(f'  {len(df):,} names with a forensic_xr_score > 0', file=sys.stderr)
    df = df.sort_values('forensic_xr_score', ascending=False).head(args.n)

    n = lambda c: pd.to_numeric(df[c], errors='coerce') if c in df.columns \
        else pd.Series(np.nan, index=df.index)
    mc = n('market_cap_usd')
    # size-normalized source components (recomputed for display)
    pension = (n('pension_funded_status').clip(lower=0) / mc)
    dta = (n('deferred_tax_valuation_allowance') / mc)
    lifo = (n('lifo_reserve') / mc)
    eqst = (n('equity_method_investments') / mc)
    rpo = (n('rpo') / n('revenue_ttm_usd'))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Forensic-XR'
    tab_colors.set_tab(ws, tab_colors.FAMILY_COLORS['deep_value'])

    f_bold = _font(bold=True, color=INK)
    f_italic_muted = _font(italic=True, color=MUTED)
    f_text = _font(color=INK)

    t = ws.cell(row=2, column=1, value='Forensic-XR — most asymmetric hidden value')
    t.font = f_bold
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOLS)
    ws.row_dimensions[2].height = 22
    sub = ws.cell(row=3, column=1,
                  value='Hidden/unpriced forensic value as a share of market cap (size-comparable), '
                        'x1.5 where the visible business is also cheap; tradeable + non-melting only.')
    sub.font = f_italic_muted
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NCOLS)
    _section_rule(ws, 5, 'Ranked by forensic_xr_score — components show the SOURCE of the hidden value',
                  span_cols=NCOLS)

    hdr = 6
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=hdr, column=i, value=h)
        c.font = _font(bold=True, color=MUTED)
        c.alignment = _TXT_ALIGN_LEFT if i in (2, 3, 4, 5, 6) else _NUM_ALIGN_CENTER
    for c in range(1, NCOLS + 1):
        ws.cell(row=hdr + 1, column=c).border = Border(top=Side(style='thin', color=INK))

    r = hdr + 1
    for rank, (_, row) in enumerate(df.iterrows(), start=1):
        r += 1
        i = row.name
        ws.cell(row=r, column=1, value=rank).font = f_text
        ws.cell(row=r, column=2, value=str(row.get('symbol') or '')).font = f_text
        ws.cell(row=r, column=3, value=str(row.get('name') or '')[:30]).font = f_text
        ws.cell(row=r, column=3).alignment = _TXT_ALIGN_LEFT
        ws.cell(row=r, column=4, value=str(row.get('src') or '')).font = _font(color=MUTED)
        ws.cell(row=r, column=5, value=str(row.get('sector') or '')[:14]).font = _font(color=MUTED)
        ws.cell(row=r, column=5).alignment = _TXT_ALIGN_LEFT
        ws.cell(row=r, column=6, value=str(row.get('industry') or '')[:22]).font = _font(color=MUTED)
        ws.cell(row=r, column=6).alignment = _TXT_ALIGN_LEFT
        _write_money(ws, r, 7, row.get('market_cap'), font=f_text)
        _v = row.get('verdict')
        _verdict_badge(ws, r, 8, _v if (pd.notna(_v) and _v) else 'UNRESEARCHED')
        _write_score(ws, r, 9, row.get('forensic_xr_score'), font=f_text)
        _write_pct(ws, r, 10, row.get('forensic_hidden_pct'), font=f_text)
        _write_pct(ws, r, 11, pension.get(i), font=f_text)
        _write_pct(ws, r, 12, dta.get(i), font=f_text)
        _write_pct(ws, r, 13, lifo.get(i), font=f_text)
        _write_pct(ws, r, 14, eqst.get(i), font=f_text)
        _write_pct(ws, r, 15, rpo.get(i), font=f_text)
        _write_score(ws, r, 16, row.get('ev_ebitda'), font=f_text)
        _write_score(ws, r, 17, row.get('pb'), font=f_text)
        _write_pct(ws, r, 18, row.get('fcf_yield'), font=f_text)
        _write_score(ws, r, 19, row.get('fq_beneish_m'), font=f_text)
        _write_pct(ws, r, 20, row.get('fq_sloan_accruals'), font=f_text)
        _write_score(ws, r, 21, row.get('fq_cfo_to_ni'), font=f_text)
        _write_score(ws, r, 22, row.get('fq_forensic_red_count'), font=f_text)
        _write_score(ws, r, 23, row.get('fq_forensic_green_count'), font=f_text)
        _ac = acct_check(row)
        c = ws.cell(row=r, column=24, value=_ac)
        c.font = _font(color=('B42318' if _ac.startswith('WARN') else INK))
        c.alignment = _TXT_ALIGN_LEFT
        for c in range(1, NCOLS + 1):
            ws.cell(row=r, column=c).border = Border(bottom=Side(style='thin', color=RULE))

    widths = {1: 4, 2: 9, 3: 30, 4: 6, 5: 14, 6: 16, 7: 13, 8: 12, 9: 11,
              10: 11, 11: 10, 12: 11, 13: 8, 14: 11, 15: 10, 16: 10, 17: 7, 18: 10,
              19: 10, 20: 11, 21: 8, 22: 6, 23: 7, 24: 44}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'A{hdr + 2}'
    if ws.max_row > hdr + 1:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(NCOLS)}{ws.max_row}"

    _accounting_sheet(wb, args.n)
    wb.save(args.out)
    print(f'wrote {args.out}: {len(df)} rows', file=sys.stderr)


def _accounting_sheet(wb, n_top: int) -> None:
    """'Accounting-Tells': the names people actually look at (top by
    entry-today asymmetry) with their quarterly forensic read, so a thesis
    resting on numbers that trip Beneish / accrual / divergence tells is
    visible next to the rank. Surfaced, never a filter."""
    df = pd.read_csv('asymmetry_global.csv', low_memory=False).drop_duplicates('symbol')
    if not os.path.exists('archetype_tags.csv'):
        return
    have = pd.read_csv('archetype_tags.csv', nrows=0).columns
    t = pd.read_csv('archetype_tags.csv', low_memory=False,
                    usecols=['symbol'] + [c for c in ACCT_COLS if c in have])
    df = df.drop(columns=[c for c in ACCT_COLS if c in df.columns], errors='ignore') \
           .merge(t.drop_duplicates('symbol'), on='symbol', how='left')
    from otc_flag import apply_otc_mode as _apply_otc
    df = _apply_otc(df, 'ex-otc')
    df = df[df.get('fmp_q_status').astype(str) == 'ok'] if 'fmp_q_status' in df.columns else df
    df['entry_today_asymmetry'] = pd.to_numeric(df.get('entry_today_asymmetry'), errors='coerce')
    df = df.sort_values('entry_today_asymmetry', ascending=False).head(n_top)
    ws = wb.create_sheet('Accounting-Tells')
    tab_colors.set_tab(ws, tab_colors.FAMILY_COLORS['deep_value'])
    hd = ['#', 'Ticker', 'Name', 'Country', 'Sector', 'Mcap (USD)', 'ETA',
          'Beneish M', 'Accruals %TA', 'CFO/NI', 'Rev streak (m)', 'Net-debt decline (m)',
          'Red', 'Green', 'Accounting check']
    ws.cell(row=2, column=1, value='Accounting tells on the top names by entry-today asymmetry').font = _font(bold=True, color=INK)
    ws.cell(row=3, column=1, value='Quarterly 3-statement forensics (FMP). Red = Beneish / accruals / '
            'receivable or inventory divergence / SBC-funded cash; green = FCF leading NI, working-capital '
            'release, proven deleveraging, gross-margin inflection, deferred-revenue build, cash-tax shield. '
            'Informational: nothing here removes a name.').font = _font(italic=True, color=MUTED)
    hdr = 5
    for i, h in enumerate(hd, start=1):
        c = ws.cell(row=hdr, column=i, value=h)
        c.font = _font(bold=True, color=MUTED)
    r = hdr
    for rank, (_, row) in enumerate(df.iterrows(), start=1):
        r += 1
        ws.cell(row=r, column=1, value=rank)
        ws.cell(row=r, column=2, value=str(row.get('symbol') or ''))
        ws.cell(row=r, column=3, value=str(row.get('name') or '')[:30])
        ws.cell(row=r, column=4, value=str(row.get('src') or ''))
        ws.cell(row=r, column=5, value=str(row.get('sector') or '')[:14])
        _write_money(ws, r, 6, row.get('market_cap_usd'))
        _write_score(ws, r, 7, row.get('entry_today_asymmetry'))
        _write_score(ws, r, 8, row.get('fq_beneish_m'))
        _write_pct(ws, r, 9, row.get('fq_sloan_accruals'))
        _write_score(ws, r, 10, row.get('fq_cfo_to_ni'))
        _write_score(ws, r, 11, row.get('fq_rev_yoy_streak_m'))
        _write_score(ws, r, 12, row.get('fq_netdebt_decline_months'))
        _write_score(ws, r, 13, row.get('fq_forensic_red_count'))
        _write_score(ws, r, 14, row.get('fq_forensic_green_count'))
        _ac = acct_check(row)
        c = ws.cell(row=r, column=15, value=_ac)
        c.font = _font(color=('B42318' if _ac.startswith('WARN') else INK))
        c.alignment = _TXT_ALIGN_LEFT
    for col, w in {1: 5, 2: 10, 3: 30, 4: 6, 5: 14, 6: 13, 7: 8, 8: 10, 9: 11, 10: 8,
                   11: 12, 12: 16, 13: 6, 14: 7, 15: 60}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'A{hdr + 1}'
    if ws.max_row > hdr:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(hd))}{ws.max_row}"


if __name__ == '__main__':
    main()
